"""Build the five speaker rehearsal script from its versioned Markdown source.

Use an isolated Python with python-docx. No model calls or evidence mutations.
The published document needs a render and visual review after rebuilding.
"""
from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/presentation/SPEAKER_SCRIPT.md"
BLACK = RGBColor(0, 0, 0)


def plain(text):
    return text.replace("**", "").replace("`", "")


def inline(p, text):
    for part in re.split(r"(\*\*.*?\*\*)", text):
        run = p.add_run(part[2:-2] if part.startswith("**") else part.replace("`", ""))
        if part.startswith("**"):
            run.bold = True


def table(doc, headings, rows, widths):
    t = doc.add_table(rows=1, cols=len(headings))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, width in zip(t.columns, widths):
        col.width = Inches(width)
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "bottom", "left", "right", "insideH", "insideV"):
        b = OxmlElement("w:"+side)
        for name, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
            b.set(qn("w:"+name), value)
        borders.append(b)
    t._tbl.tblPr.append(borders)
    for i, values in enumerate([headings, *rows]):
        cells = t.rows[0].cells if i == 0 else t.add_row().cells
        rowprops = cells[0]._tc.getparent().get_or_add_trPr()
        rowprops.append(OxmlElement("w:cantSplit"))
        if i == 0:
            rowprops.append(OxmlElement("w:tblHeader"))
        for j, (c, value) in enumerate(zip(cells, values)):
            c.width = Inches(widths[j])
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = c._tc.get_or_add_tcPr()
            margins = OxmlElement("w:tcMar")
            for side in ("top", "left", "bottom", "right"):
                m = OxmlElement("w:"+side)
                m.set(qn("w:w"), "100")
                m.set(qn("w:type"), "dxa")
                margins.append(m)
            tcpr.append(margins)
            if i == 0:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "E7EEF4")
                tcpr.append(shading)
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.08
            if j in (1, 3):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(plain(value))
            r.font.size = Pt(10.5)
            r.bold = i == 0
    doc.add_paragraph()


def build(output):
    markdown = SOURCE.read_text(encoding="utf-8")
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.top_margin = s.bottom_margin = Cm(2)
    s.left_margin = s.right_margin = Cm(2.2)
    s.header_distance = s.footer_distance = Cm(0.8)
    styles = doc.styles
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "Footer", "Header", "Quote"):
        styles[name].font.name = "Calibri"
        styles[name].font.color.rgb = BLACK
        styles[name].element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Calibri")
        for color in styles[name].element.xpath(".//w:color"):
            for attribute in ("themeColor", "themeTint", "themeShade"):
                color.attrib.pop(qn("w:"+attribute), None)
        for border in styles[name].element.xpath(".//w:pBdr"):
            border.getparent().remove(border)
    normal = styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.widow_control = True
    styles["Title"].font.size = Pt(26)
    styles["Title"].font.bold = True
    for name, size in (("Heading 1", 19), ("Heading 2", 13.5), ("Heading 3", 12)):
        style = styles[name]
        style.font.size = Pt(size)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(13)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.keep_with_next = True
    h = s.header.paragraphs[0]
    h.add_run("EDITGUARD  PRESENTATION SCRIPT").font.size = Pt(8)
    f = s.footer.paragraphs[0]
    f.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    f.add_run("Page ").font.size = Pt(8)
    r = f.add_run()
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    r._r.addnext(field)
    doc.core_properties.title = "EditGuard Presentation Script"
    doc.core_properties.author = "EditGuard project team"
    doc.core_properties.subject = "Five speaker rehearsal and contribution allocation"
    doc.core_properties.keywords = "EditGuard presentation script speakers demonstration contributions"

    lines = markdown.splitlines()
    i = 0
    member_started = False
    contribution = False
    while i < len(lines):
        line = lines[i].strip()
        i += 1
        if not line:
            continue
        if line.startswith("# "):
            doc.add_paragraph("EditGuard Presentation Script", "Title")
        elif line.startswith("## "):
            heading = line[3:]
            if heading.startswith("Member ") or heading in ("Optional live demonstration", "Short answers for panel questions", "Contribution allocation for the five members"):
                p = doc.add_heading(heading, level=1)
                p.paragraph_format.page_break_before = True
                p.paragraph_format.space_before = Pt(0)
            else:
                doc.add_heading(heading, level=1)
            contribution = heading == "Contribution allocation for the five members"
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=2)
        elif line.startswith("| "):
            table_lines = [line]
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            parsed = [[s.strip() for s in row.strip("|").split("|")] for row in table_lines]
            headings, rows = parsed[0], parsed[2:]
            if contribution:
                for values in rows:
                    p = doc.add_heading(plain(values[0])+" "+plain(values[1]), level=2)
                    p.paragraph_format.keep_with_next = True
                    p = doc.add_paragraph()
                    inline(p, "**Contribution:** "+values[2])
                    p.paragraph_format.keep_with_next = True
                    p = doc.add_paragraph()
                    inline(p, "**Evidence and collaboration:** "+values[3])
                    p.paragraph_format.space_after = Pt(13)
            else:
                table(doc, headings, rows, [0.8, 0.7, 3.18, 1.85])
        elif line.startswith("> "):
            p = doc.add_paragraph(line[2:])
            p.paragraph_format.line_spacing = 1.18
            p.paragraph_format.space_after = Pt(12)
            for run in p.runs:
                run.font.size = Pt(12)
        elif re.match(r"^\*\*[^*]+\*\*$", line):
            p = doc.add_paragraph()
            inline(p, line)
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_before = Pt(8)
        else:
            p = doc.add_paragraph()
            inline(p, line)
            if line.startswith("**Time:"):
                p.paragraph_format.keep_with_next = True
                for run in p.runs:
                    run.font.size = Pt(10)
            if line.startswith("**Handover:"):
                p.paragraph_format.space_before = Pt(6)
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    digest = lambda p: hashlib.sha256(p.read_bytes()).hexdigest()
    print(json.dumps({"output":str(output), "source":str(SOURCE.relative_to(ROOT)),
        "source_sha256":digest(SOURCE), "docx_sha256_before_field_refresh":digest(output),
        "spoken_words":sum(len(x.split()) for x in re.findall(r"^> (.+)$", markdown, re.M)),
        "model_calls":0}, indent=2))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output",type=Path,default=ROOT/"output/presentation/EditGuard_Speaker_Script.docx")
    build(parser.parse_args().output)
