"""Create the complete EditGuard report from released evidence without model calls.

Run with the Codex bundled Python containing python-docx and Pillow. Scientific
figures are the already audited Matplotlib exports; the architecture schematic
is generated separately and is never counted as an experimental measurement.
"""
from __future__ import annotations

import argparse
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
OPS = ROOT / "reports/reasonke/operations-final"
DEPTH = ROOT / "reports/reasonke/depth-final"
BLACK = RGBColor(0, 0, 0)


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def read(path):
    return json.loads(Path(path).read_text(encoding="utf-8-sig"))


def evidence():
    from editguard_chat.results import depth_results
    from scripts.audit_operations_evidence import audit_receipts
    from scripts.audit_india_http import audit
    from editguard_depth.metrics import audit_and_aggregate
    audit_receipts([OPS / f"final-{m}.json" for m in ("reasonke", "instruction")],
                   protocol=OPS / "protocol.json")
    audited_depth = audit_and_aggregate([DEPTH / f"final-{m}.json" for m in ("reasonke", "instruction")],
                                       protocol=DEPTH / "protocol.json")
    released = depth_results(ROOT)
    if not released["available"] or audited_depth["execution_status"] != "completed":
        raise ValueError("Complete released depth evidence required")
    audit(DEPTH / "india-acceptance.json")
    original, depth = read(OPS / "results.json"), read(DEPTH / "results.json")
    if original["assigned_records"] != 1312 or depth["assigned_records"] != 544:
        raise ValueError("Unexpected study populations")
    return original, depth


def architecture(path):
    """A schematic of implemented control flow, not a measured scientific graph."""
    im = Image.new("RGB", (1800, 1040), "white")
    d = ImageDraw.Draw(im)
    regular = ImageFont.truetype("C:/Windows/Fonts/calibri.ttf", 30)
    bold = ImageFont.truetype("C:/Windows/Fonts/calibrib.ttf", 32)
    small = ImageFont.truetype("C:/Windows/Fonts/calibri.ttf", 27)

    def box(x, y, w, h, title, lines, fill="#edf2f5"):
        d.rounded_rectangle((x, y, x+w, y+h), radius=10, fill=fill, outline="#9aaab1", width=2)
        d.text((x+22, y+18), title, font=bold, fill="black")
        for i, line in enumerate(lines):
            d.text((x+22, y+66+i*35), line, font=regular, fill="black")

    def arrow(a, b):
        d.line((a, b), fill="#356a68", width=5)
        x, y = b
        ax, ay = a
        if abs(x-ax) > abs(y-ay):
            sign = 1 if x > ax else -1
            d.polygon([(x,y),(x-sign*17,y-11),(x-sign*17,y+11)], fill="#356a68")
        else:
            sign = 1 if y > ay else -1
            d.polygon([(x,y),(x-11,y-sign*17),(x+11,y-sign*17)], fill="#356a68")

    box(20, 35, 500, 175, "Browser application", ["Chat and approved corrections", "History comparison and results"])
    box(650, 35, 500, 175, "Application controller", ["Request validation and lock", "Revision safe conversation history"])
    box(1280, 35, 500, 175, "SQLite fact memory", ["Versions sources and canonical IDs", "Atomic approval and rollback"])
    arrow((520,122),(650,122)); arrow((1150,122),(1280,122))
    box(650, 315, 500, 175, "Approved context selector", ["Two lexical seeds and expansion", "Four supplied facts in depth mode"])
    arrow((900,210),(900,315)); arrow((1510,210),(1120,315))
    box(650, 595, 500, 175, "Private neural worker", ["Pinned ReasonKE checkpoint", "Frozen weights and measured output"], fill="#e6f1ed")
    box(20, 595, 500, 175, "Answer and evidence", ["Generated answer and raw output", "Fact IDs sources and revisions"])
    arrow((900,490),(900,595)); arrow((650,680),(520,680))
    d.line([(270,595),(270,260),(120,260),(120,210)], fill="#356a68", width=4)
    d.polygon([(120,210),(109,227),(131,227)], fill="#356a68")
    d.text((25,825), "Separate offline evaluation path", font=bold, fill="black")
    d.text((25,875), "Frozen questions and protocol  ->  finite inference  ->  raw ledgers  ->  CPU replay audit  ->  tables and graphs", font=small, fill="black")
    d.text((25,945), "Approval changes external memory. The model generates the answer. Retrieved provenance does not verify every generated sentence.", font=small, fill="black")
    im.save(path)


class Report:
    def __init__(self, output):
        self.output = Path(output)
        self.doc = Document()
        s = self.doc.sections[0]
        s.page_width, s.page_height = Cm(21), Cm(29.7)
        s.top_margin, s.bottom_margin = Cm(2.05), Cm(2.0)
        s.left_margin, s.right_margin = Cm(2.2), Cm(2.2)
        s.header_distance, s.footer_distance = Cm(.8), Cm(.8)
        s.different_first_page_header_footer = True
        self.width = 6.535
        styles = self.doc.styles
        for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "Caption", "Header", "Footer"):
            st = styles[name]
            st.font.name, st.font.color.rgb = "Calibri", BLACK
            st.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Calibri")
            for border in list(st.element.xpath(".//w:pBdr")):
                border.getparent().remove(border)
        normal = styles["Normal"]
        normal.font.size = Pt(11)
        normal.paragraph_format.line_spacing = 1.12
        normal.paragraph_format.space_after = Pt(7)
        normal.paragraph_format.widow_control = True
        for name, size, before, after in (("Title", 30, 0, 14), ("Subtitle", 17, 0, 14),
                ("Heading 1", 18, 0, 13), ("Heading 2", 12.5, 12, 6), ("Heading 3", 11, 9, 4)):
            styles[name].font.size = Pt(size)
            styles[name].font.bold = name != "Subtitle"
            styles[name].paragraph_format.space_before = Pt(before)
            styles[name].paragraph_format.space_after = Pt(after)
            styles[name].paragraph_format.keep_with_next = True
        styles["Caption"].font.size = Pt(9)
        styles["Caption"].font.italic = True
        styles["Caption"].paragraph_format.space_after = Pt(9)
        for name in ("Header", "Footer"):
            styles[name].font.size = Pt(8)
        h = s.header.paragraphs[0]
        h.text = "EDITGUARD   TECHNICAL PROJECT REPORT"
        f = s.footer.paragraphs[0]
        f.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        f.add_run("EditGuard   |   Page ")
        self.field(f, "PAGE")
        settings = self.doc.settings.element
        up = OxmlElement("w:updateFields"); up.set(qn("w:val"), "true"); settings.append(up)
        self.doc.core_properties.title = "EditGuard Knowledge Editing and Verification Platform"
        self.doc.core_properties.subject = "Complete technical report and controlled experimental findings"
        self.doc.core_properties.author = "EditGuard project team"
        self.doc.core_properties.keywords = "knowledge editing, external memory, retrieval, ReasonKE, evaluation, rollback"
        self.figures = []
        self.tables = 0

    @staticmethod
    def field(paragraph, code):
        run = paragraph.add_run()
        begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
        instruction = OxmlElement("w:instrText"); instruction.set(qn("xml:space"), "preserve"); instruction.text = " " + code + " "
        separate = OxmlElement("w:fldChar"); separate.set(qn("w:fldCharType"), "separate")
        label = OxmlElement("w:t"); label.text = "1"
        end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
        for element in (begin, instruction, separate, label, end): run._r.append(element)

    def p(self, text, *, bold=False):
        p = self.doc.add_paragraph()
        p.add_run(text).bold = bold
        return p

    def prose(self, text):
        for paragraph in re.split(r"\n\s*\n", text.strip()):
            self.p(" ".join(paragraph.split()))

    def h2(self, text):
        self.doc.add_heading(text, level=2)

    def chapter(self, text):
        p = self.doc.add_heading(text, level=1)
        starts_page = text in {"Abstract", "Contents", "1 Problem and objectives",
            "8 Experimental results", "References", "Appendix A Reproduction and evidence map"}
        p.paragraph_format.page_break_before = starts_page
        p.paragraph_format.space_before = Pt(0 if starts_page else 19)
        return p

    def table(self, caption, headings, rows, widths):
        self.tables += 1
        cap = self.doc.add_paragraph(f"Table {self.tables}  {caption}", "Caption")
        cap.paragraph_format.keep_with_next = True
        table = self.doc.add_table(rows=1, cols=len(headings))
        table.alignment, table.autofit = WD_TABLE_ALIGNMENT.CENTER, False
        for col, width in zip(table.columns, widths): col.width = Inches(width)
        props = table._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            edge = OxmlElement("w:"+side)
            for key, value in (("val","single"),("sz","4"),("color","D9D9D9")):
                edge.set(qn("w:"+key), value)
            borders.append(edge)
        props.append(borders)
        for idx, values in enumerate([headings, *rows]):
            cells = table.rows[0].cells if idx == 0 else table.add_row().cells
            trpr = cells[0]._tc.getparent().get_or_add_trPr()
            keep = OxmlElement("w:cantSplit"); trpr.append(keep)
            if idx == 0:
                repeat = OxmlElement("w:tblHeader"); trpr.append(repeat)
            for j, (cell, value, width) in enumerate(zip(cells, values, widths)):
                cell.width, cell.vertical_alignment = Inches(width), WD_CELL_VERTICAL_ALIGNMENT.CENTER
                tcpr = cell._tc.get_or_add_tcPr()
                margins = OxmlElement("w:tcMar")
                for side in ("top", "left", "bottom", "right"):
                    item = OxmlElement("w:"+side); item.set(qn("w:w"),"85"); item.set(qn("w:type"),"dxa"); margins.append(item)
                tcpr.append(margins)
                if idx == 0:
                    shade = OxmlElement("w:shd"); shade.set(qn("w:fill"),"E7EEF4"); tcpr.append(shade)
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.line_spacing = 1.03
                # Keep short comparison tables together and never strand a
                # repeating table header below its caption without a data row.
                p.paragraph_format.keep_with_next = (idx < len(rows) if len(rows) <= 6 else idx == 0)
                if j and len(str(value)) < 24: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = p.add_run(str(value)); r.font.size = Pt(9.3); r.bold = idx == 0; r.font.color.rgb = BLACK
        self.doc.add_paragraph().paragraph_format.space_after = Pt(0)

    def figure(self, path, caption, width=None):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run()
        picture = run.add_picture(str(path), width=Inches(width or self.width))
        picture._inline.docPr.set("descr", caption)
        number = len(self.figures)+1
        self.doc.add_paragraph(f"Figure {number}  {caption}", "Caption")
        self.figures.append({"number": number, "path": str(Path(path).relative_to(ROOT)), "sha256": sha(path), "caption": caption})

    def code(self, lines):
        for line in lines.strip().splitlines():
            p = self.doc.add_paragraph()
            p.paragraph_format.space_after = Pt(3)
            r = p.add_run(line); r.font.name = "Consolas"; r.font.size = Pt(8.7)


def build(output):
    original, depth = evidence()
    report = Report(output)
    assets = ROOT / "output/docx/assets"
    assets.mkdir(parents=True, exist_ok=True)
    diagram = assets / "editguard-architecture.png"
    architecture(diagram)
    d = report.doc
    p = d.add_paragraph(); p.paragraph_format.space_after = Pt(50)
    d.add_paragraph("EditGuard Knowledge Editing and Verification Platform", "Title")
    d.add_paragraph("Technical Project Report", "Subtitle")
    report.p("Governed external memory and measured factual reasoning for conversational AI")
    report.p("EditGuard project team\n29 October 2026")
    report.prose("""
    This report presents the implemented platform, its independently authored controlled datasets,
    two frozen comparative evaluations, and the live application checks. The system makes an approved
    factual correction visible, scoped, traceable and reversible while an upstream trained model
    generates answers from the applicable memory.

    The principal finding is that selected approved context supported perfect observed short answer
    scores for ReasonKE on both controlled main tests, with a substantial latency cost. Full memory
    produced errors, the comparison model retained abstentions, and correct final answers did not
    certify every generated explanation. These boundaries are part of the result.
    """)
    report.p("Completed evidence  1312 original study requests  544 depth study requests  14 India application generations", bold=True)
    report.p("Project repository  github.com/ShivamRaghav7/verity\nRelease branch  codex/reasonke-integration")

    report.chapter("Abstract")
    report.prose("""
    A conversational assistant can use an obsolete service owner even after an organisation updates
    its records. The resulting error may propagate: the wrong owner implies the wrong on call contact
    and the wrong contact location. Updating the underlying information is necessary, but an assistant
    must also select the right supporting facts, keep historical releases distinct, and preserve
    unrelated answers. EditGuard addresses this bounded maintenance problem through explicit
    approval, versioned external memory, context selection, neural answer generation and replayable evaluation.

    The final application integrates the author published ReasonKE_plus checkpoint with a lightweight
    Python website, SQLite history, a private inference worker and an evidence dashboard. New corrections
    update memory without modifying model weights. The project contributes the governed maintenance
    workflow, retrieval adaptations, independent fictional datasets, matched comparisons, finite execution
    and recovery, raw output audits and a deployable interface. Foundation model training and the
    ReasonKE research method remain attributed to their upstream authors.

    The original operations evaluation contains 1312 generations: five matched main configurations on
    256 questions over 32 held out lineages, plus two 16 question oracle diagnostics. ReasonKE with
    updated retrieved facts scored 256/256; Qwen with the same updated retrieval scored 241/256 with
    fifteen abstentions. ReasonKE with full memory scored 231/256, below Qwen full memory at 241/256.
    A separate depth evaluation contains 544 generations, including 64 three hop questions per main
    model condition. ReasonKE scored 256/256 main answers and Qwen 248/256, with eight abstentions.
    Three hop scores were 64/64 and 63/64 respectively. Both models scored 16/16 on each study's
    separate preselected gold support diagnostic.

    The depth main difference was 3.125 percentage points, with a 95 percent paired lineage bootstrap
    interval from 1.5625 to 5.078125 points. ReasonKE's warm median and 95th percentile latency were
    15.94 and 19.28 seconds, compared with Qwen's 1.12 and 1.38 seconds. Fourteen actual India website
    answers and twelve consistency checks passed, including conflict rejection, historical comparison
    and atomic rollback. Labels were audited automatically, without independent human semantic review.
    The perfect observed scores concern small explicit fact worlds and repeated within split templates;
    they establish neither universal question accuracy nor production reliability.
    """)

    toc_heading = d.add_paragraph("Contents", "Normal")
    toc_heading.paragraph_format.page_break_before = True
    toc_heading.paragraph_format.space_after = Pt(13)
    toc_heading.runs[0].font.size = Pt(18)
    toc_heading.runs[0].bold = True
    p = d.add_paragraph()
    report.field(p, 'TOC \\o "1-1" \\h \\z \\u')
    report.p("The main chapters describe the current text platform. Earlier auxiliary model and visual studies remain archived and are not pooled into these results.")

    report.chapter("1 Problem and objectives")
    report.h2("Application problem")
    report.prose("""
    Operational knowledge changes in connected records. A service moves to a new owning team, while
    each team's contact and each person's office remain separate facts. A user may ask for the owner,
    the appropriate contact, or the office of that contact using several descriptions of the service.
    A correction to ownership should change the affected reasoning path while leaving the service's
    hosting location, another product tier and an older release unchanged.

    A web page or structured lookup is often sufficient for a handful of published policies. The
    stronger use case here is interactive querying across related operational records, with follow up
    questions, explicit evidence and historical comparisons. EditGuard is a controlled demonstration
    of that use case, not evidence of commercial demand or a deployed enterprise incident service.
    We deliberately keep short factual tasks separate from unrestricted dialogue.
    """)
    report.h2("Concrete maintenance scenario")
    report.prose("""
    In the fictional India handbook, Kaveri Payment Gateway in release 2 initially belongs to Bengaluru
    Platform Team. That team's contact is Meera Iyer, whose office is in Chennai. An administrator can
    replace the owner with Pune Reliability Team. Its contact is Arjun Deshmukh, whose office is in
    Hyderabad. The hosting location remains Bengaluru. Release 1 and the Pro service retain their
    own owners. Indian place names make the demonstration familiar; the service, people, teams and
    assignments are invented.

    The three hop answer therefore depends on a chain of three relations, rather than on the city in
    a team's name. Before the ownership update the expected office answer is Chennai; after it the
    expected answer is Hyderabad. Rollback should restore the original path. This is the intended
    behaviour; only recorded model responses establish whether a generation achieved it.
    """)
    report.h2("Objectives and acceptance criteria")
    report.table("Project objectives and observable acceptance", ["Objective", "Acceptance evidence"], [
        ["Govern corrections", "Only an explicit operator approval makes a fact authoritative"],
        ["Respect scope and history", "Separate canonical entity property and release keys plus historical revision comparison"],
        ["Support connected answers", "Direct paraphrase two hop and three hop questions scored from raw neural outputs"],
        ["Make changes reversible", "Atomic replacement and rollback with retained versions and event history"],
        ["Measure tradeoffs fairly", "Matched updated contexts output budgets latency and all assigned outcomes"],
        ["Make results inspectable", "Frozen protocols raw ledgers CPU replay scientific exports and a runnable website"]
    ], [2.05, 4.48])
    report.prose("""
    Completion is defined by a working product and reproducible evidence, not by a guaranteed accuracy
    target. Deterministic transaction correctness and generated answer correctness are tested
    separately. A successful database restore does not count as a neural answer, and an operator
    supplied source does not independently prove that a statement is true.
    """)

    report.chapter("2 Background and method choice")
    report.h2("What knowledge editing means in this project")
    report.prose("""
    Knowledge editing aims to change a model system's response to selected information while preserving
    behaviour outside the intended scope. Parametric approaches alter model weights. Memory based
    approaches retain corrections outside the base model and condition behaviour on retrieved edits.
    EditGuard implements the latter operational route: its learned generator was trained upstream,
    while local approval changes external memory. The selector and transaction rules are deterministic
    software. The current deployment has no newly trained standalone scope classifier.
    """)
    report.h2("Relevant research")
    report.prose("""
    Retrieval augmented generation combines a language model with external retrievable information.
    Lewis and colleagues studied this approach for knowledge intensive tasks [1]. EditGuard uses
    retrieval as a practical way to supply current facts, with explicit version and approval control.
    Our lexical selector and operational schema are adaptations; they are not a reproduction of the
    original RAG training architecture.

    SERAC stores edits in explicit memory and learns their applicability and counterfactual response
    behaviour [2]. This motivated the initial separation of edit memory, scope and answer generation.
    The final ReasonKE application does not contain SERAC's trained auxiliary components and is not
    described as a faithful SERAC implementation.

    ROME demonstrates direct changes to factual associations by editing feed forward weights [3].
    It provides a useful contrast: EditGuard does not perform this kind of per correction weight edit.
    Its historical states refer to approved memory revisions rather than alternative neural checkpoints.

    ReasonKE research uses explicit reasoning chains to guide the use of new facts in multi hop
    questions [4]. ReasonKE++ studies process level faithfulness through supervised and reinforcement
    learning [5]. We integrate an author published ReasonKE_plus checkpoint [6], but its exact
    correspondence to the paper's reported checkpoint is undocumented. We neither repeat its training
    nor substitute a published benchmark number for our own measurements.
    """)
    report.table("Method families and the implemented choice", ["Family", "What changes", "Use in this report"], [
        ["Parametric editing", "Selected model weights", "Background comparison only"],
        ["SERAC style editing", "Memory plus trained edit routing and response modules", "Earlier motivation not the final architecture"],
        ["Updated retrieval", "Approved external information and selected context", "Both models receive the same current information"],
        ["ReasonKE integration", "Memory supplied to an upstream trained generator", "Final application and measured comparisons"]
    ], [1.42, 2.5, 2.61])
    report.h2("Why this implementation was selected")
    report.prose("""
    Reversibility and provenance are central to the product. External versions make an edit inspectable
    and allow exact historical reconstruction without storing another set of model weights. A compact
    serving layer keeps the GPU process separate from the website. The completed development
    comparisons selected retrieved context before final evaluation. The choice is supported by those
    measurements; it is not a claim that retrieval always beats weight editing or that ReasonKE is the
    best method on every knowledge editing task.
    """)

    report.chapter("3 System architecture and inference")
    report.prose("""
    The online application and offline experiment pipeline share model and retrieval interfaces but have
    different responsibilities. Online requests read approved memory and produce conversational
    responses. Offline runs use a fixed question plan, save every assigned outcome, and produce
    audited measurements. The dashboard displays saved evidence rather than silently evaluating the
    current conversation.
    """)
    report.figure(diagram, "Implemented online data flow and separate evaluation path. This is an architecture schematic, not a performance measurement.")
    report.h2("Online request flow")
    report.prose("""
    The browser sends a text question to the application controller. The controller validates the
    request, serializes memory and generation operations, and reads the active approved facts at a
    specific knowledge revision. It selects supporting facts, adds permitted conversation history and
    calls the private neural worker. The worker generates text using its pinned checkpoint. The
    controller stores the complete user and assistant pair, raw output, supplied fact IDs, sources,
    generation statistics and model and knowledge revisions before returning the response.

    Conversation history is revision aware. Complete user and assistant pairs from the current
    revision can enter a new prompt. Older pairs remain visible but are excluded from new prompts,
    reducing contamination by answers generated before an edit. Historical comparisons use matched
    pre question history and exclude correction conditioned pairs from both sides.
    """)
    report.h2("Context selection")
    report.prose("""
    The original operations study uses a two stage lexical selector. A BM25 style ranker scores fact
    text against the question, takes two seed facts, adds their answer text to the query, then fills a
    four fact context after reranking. Tie breaks use stable IDs. Normal retrieval does not read the
    expected answer or gold supporting fact labels.

    The separate depth study and India launcher use bounded iterative expansion. Two question ranked
    seeds initialize an answer text frontier. The selector adds a highest ranked unseen fact from
    that frontier and advances to its answer text, with at most three rounds and four total facts.
    Development selection considered a declared family of seed counts and context caps before final
    inference. The default CLI still uses the original policy; the India launcher explicitly requests
    iterative mode. The two policies must not be described as one unchanged experiment.

    Both selectors are bounded lexical software rather than trained applicability classifiers. Scope
    identity and retrieval can reduce mistakes, but the implementation does not supply a calibrated
    probability that a correction applies to an arbitrary natural language question.
    """)

    report.chapter("4 Knowledge lifecycle and correctness")
    report.h2("Stored facts and identity")
    report.table("Core approved fact fields", ["Field group", "Recorded information"], [
        ["Fact content", "Question or relation description answer scope and required source"],
        ["Canonical identity", "Entity ID relation ID and scope ID supplied together optional aliases"],
        ["Version lineage", "Fact UUID lineage ID version predecessor and successor links"],
        ["Lifecycle state", "Active superseded or withdrawn status and approval metadata"],
        ["Audit history", "Actor timestamps action and knowledge revision for each event"]
    ], [1.58, 4.95])
    report.prose("""
    The canonical key is the combination of entity ID, relation ID and scope ID. Separate release or
    plan IDs allow multiple applicable histories without collapsing their facts. Where canonical IDs
    are absent, the legacy key uses normalized question and scope text. SQLite enforces at most one
    active version per key. Aliases are metadata, not evidence of a learned general entity linker.

    Required text fields are bounded and validated. A source describes the operator's supplied
    provenance; it is not an independent fact checking service. Version timestamps record when a
    change was approved. The application does not implement a general calendar effective date or
    overlapping date interval policy engine. Its historical interface is based on knowledge revisions.
    """)
    report.h2("Approval replacement and conflicts")
    report.prose("""
    A customer message never becomes an authoritative correction automatically. The operator must
    explicitly approve a fact. A competing approval for an already active key is rejected with HTTP
    409. To replace it, the operator names the active predecessor through the supersedes field. One
    transaction creates the new active version, supersedes the old version, writes audit events and
    increments the knowledge revision. A failure rolls back the entire change.

    The deterministic guarantee concerns the represented key. Two semantically contradictory statements
    written under unrelated keys may evade this check. The system is therefore a governed local
    maintenance workflow, not a general natural language contradiction detector. Ambiguous intent
    and unrepresented scope still require operator judgement or a further question.
    """)
    report.h2("Rollback and historical comparison")
    report.prose("""
    Atomic rollback applies to an active replacement with a direct predecessor of the same key. In a
    single BEGIN IMMEDIATE transaction it withdraws the replacement, restores the predecessor and
    records both events at one revision. It refuses an obsolete replacement that would undo newer
    state. Repeating an already completed rollback is idempotent in the permitted restored state.

    Historical reads replay ordered events up to a requested valid revision without changing current
    memory. A before and after comparison generates two fresh answers from the same model, with
    different approved snapshots. The alternative without correction context uses empty edit memory;
    this is different from the approved pre edit snapshot. Context is reread after changes and history
    is revision filtered. The ReasonKE selector has no persistent generated answer cache to preserve
    a stale answer across the update.
    """)

    report.chapter("5 Data and explicit ground truth")
    report.h2("Why controlled fictional data")
    report.prose("""
    A fictional catalogue lets us define the original fact, corrected fact, unchanged relations and
    historical alternatives precisely. A model cannot be assumed to know an invented company from
    pretraining. Every scored request receives its configured initial or updated facts explicitly.
    This makes the consequences of a specified edit testable, while limiting ecological validity.

    The data are independently authored for this project rather than copied from an upstream benchmark.
    They are controlled demonstration data, with automatically audited graph paths and labels. There
    was no independent human semantic review of the final wording. Automatic consistency checks help
    detect malformed records, but cannot establish that every question is natural or unambiguous.
    """)
    report.table("Authored packets and measured populations", ["Quantity", "Original operations", "Separate depth"], [
        ["Authored graph lineages", "48", "48"], ["Facts per world", "9", "12"],
        ["Authored questions", "384", "384"], ["Development lineages and questions", "16 and 128", "16 and 128"],
        ["Final lineages and questions", "32 and 256", "32 and 256"],
        ["Measured neural development calls", "120 across five conditions", "48 across two models"],
        ["Final main conditions", "5", "2"], ["Final assigned generations including oracle", "1312", "544"]
    ], [2.56, 1.98, 1.99])
    report.h2("Split and wording controls")
    report.prose("""
    Lineages are partitioned before question augmentation so related original and edited variants stay
    together. Development guides model and retrieval choices; the frozen final set supplies the reported
    held out measurement. The original packet repeats wording templates across the split and tests new
    assignments more strongly than new language. The depth packet reserves distinct authored
    development and final wording banks, with disjoint service, team and person identities across the
    split. It still reuses templates within each split, city values, relation vocabulary and canonical
    fact syntax. Unseen entities therefore do not imply unrestricted linguistic generalization.

    The authored 128 question development partition is larger than the finite neural development
    workload. The depth profile uses eight requests per model and development uses 24 per model.
    These are declared subsets, not evidence that all authored development questions were generated
    under every configuration. The final questions and scoring were fixed before inference.
    """)
    report.h2("Depth and locality categories")
    report.table("Depth study question groups per main model", ["Group", "Relation path", "Assigned"], [
        ["Locality controls", "Other property or historical release", "64"],
        ["One hop", "Service to owning team", "64"],
        ["Two hops", "Service to team to contact", "64"],
        ["Three hops", "Service to team to contact to office city", "64"]
    ], [1.38, 4.3, .85])
    report.prose("""
    Direct and paraphrased forms each contribute 32 questions to every applicable depth. The two
    locality categories contribute 32 each. A depth zero label denotes a control, not a task requiring
    no evidence: it still needs an independent fact. Wrong releases, similar products and other
    relations remain explicit alternatives in the memory. The display handbook is another separate
    fictional example and does not rename final test records after evaluation.
    """)

    report.chapter("6 Experimental design and scoring")
    report.h2("Matched configurations")
    report.table("Original main conditions", ["Label", "Generator", "Provided information"], [
        ["R0", "ReasonKE", "Initial facts scored against updated truth staleness diagnostic"],
        ["R RAG", "ReasonKE", "Updated approved facts selected by retrieval"],
        ["R ALL", "ReasonKE", "The full updated approved memory"],
        ["G RAG", "Qwen instruction model", "The same updated selection policy"],
        ["G ALL", "Qwen instruction model", "The same full updated approved memory"]
    ], [.75, 1.65, 4.13])
    report.prose("""
    Updated baselines receive the same approved corrections and supporting records. R0 intentionally
    represents stale information and is a diagnostic reference, not a claim that a properly updated
    RAG assistant was left stale. Full memory is the prompt based correction comparison. The final
    matrix does not include every editing method or a newly trained scope classifier. Any conclusion
    must stay within these implemented comparisons.

    The depth extension compares the two checkpoints using the same selected iterative policy. Each
    model has 256 main requests and sixteen preselected three hop gold support requests. An oracle
    diagnostic supplies the required facts directly to isolate answer generation from ordinary
    retrieval. It remains a distinct population; its successes never replace main failures.
    """)
    report.h2("Metrics and denominators")
    report.prose("""
    Primary factual accuracy is the number of correct final answers divided by all assigned questions.
    Wrong answers, abstentions, execution errors and unexecuted requests receive zero primary credit.
    Coverage is the fraction producing an endorsed answer rather than an abstention or execution
    failure. Answered only precision conditions on that answered population and is reported separately.
    An UNKNOWN output is a generated abstention, not an absent request.

    Retrieval required fact recall counts supplied gold required fact occurrences divided by all required
    occurrences. Complete path recall counts questions for which every required supporting fact was
    supplied. These measures audit routing; they do not imply that the generator answered correctly.
    Ordinary retrieval does not receive gold IDs. Gold IDs are used by the evaluator after selection
    and by the separately labelled oracle condition.

    Baseline correct retention measures how many originally correct out of scope answers remain
    correct. Overall locality accuracy measures correctness on every locality question, including those
    the initial baseline missed. False correction application specifically counts endorsement of an
    inapplicable edited value. It is narrower than the set of all wrong answers. The depth study has no
    unchanged baseline condition, so it supplies no new baseline retention estimate.
    """)
    report.h2("Answer matching and uncertainty")
    report.prose("""
    The scorer compares the extracted endorsed final answer with frozen reviewed canonical aliases
    attached to a question's subject and relation. A target merely mentioned as an example or rejected
    possibility is insufficient. Strict response contracts can reject an answer that contains correct
    information but does not provide the requested short span. Raw outputs remain available for
    interpretation. No sole LLM judge determines these scores, and no template answer substitutes for
    a failed neural generation.

    Paired comparisons reuse the same questions and resample whole lineages, with 10000 bootstrap
    resamples and seed 42. Eight question variants from one world are correlated; treating 256
    variants as 256 independent natural language tasks would overstate precision. Intervals describe
    the constructed population and do not remove uncertainty about real deployment traffic. A zero
    width interval from identical sampled outcomes does not establish universal certainty.
    """)

    report.chapter("7 Model configuration and finite execution")
    report.table("Fixed final inference configuration", ["Setting", "Value"], [
        ["ReasonKE checkpoint", "YukinoKaren/ReasonKE_plus"],
        ["Comparison checkpoint", "Qwen/Qwen2.5-7B-Instruct"],
        ["Parameters per checkpoint", "7615616512 approximately 7.62 billion"],
        ["Quantization and computation", "NF4 double quantization with FP16 computation"],
        ["Decoding", "Greedy one beam no sampling"],
        ["Input and new token limits", "4096 input tokens and at most 512 new tokens"],
        ["Normal selected context", "At most four approved facts"],
        ["GPU execution", "Tesla T4 one checkpoint loaded at a time"],
        ["Measured depth software", "Python 3.13.16 Torch 2.11.0+cu130 CUDA 13.0 Transformers 4.56.1"],
        ["Auxiliary runtime packages", "Accelerate 1.10.1 bitsandbytes 0.48.1"]
    ], [2.05, 4.48])
    report.prose("""
    Both checkpoints were trained upstream. Qwen's official model card describes its instruction tuned
    transformer architecture [7]. EditGuard loads pinned safe tensor revisions without remote model
    code, measures inference resources and conditions answers on approved facts. We performed no
    new checkpoint fine tuning or per correction gradient update for these final studies. Training
    and validation loss curves would therefore be misleading in this report. Earlier auxiliary model
    experiments are preserved independently and do not become evidence of training this final model.

    Profiling preceded substantial inference. The depth eight request profiles measured mean times
    of 15.97 seconds for ReasonKE and 1.09 seconds for Qwen. Their uncertain extrapolations excluded
    downloads and loading. Final runs had finite assigned populations, 1800 second load watchdogs,
    240 second request watchdogs and an eight hour per model cap. These are stopping limits, not
    advertised runtimes. Peak depth loading allocation was 6.59 GiB per model and peak generation
    allocation approximately 5.39 GiB. Allocation and reservation are not total GPU capacity.

    Every request outcome is written atomically. Exact plan resume verifies source, data, settings and
    protocol identity and skips recorded successes, wrong answers, abstentions and errors. It runs only
    unfinished assignments. The final protocols were frozen before final generations; no test outcome
    justified prompt, alias or threshold changes.
    """)

    report.chapter("8 Experimental results")
    report.h2("Original operations comparison")
    labels = {"R0":"ReasonKE initial facts", "R-RAG":"ReasonKE retrieved", "R-ALL":"ReasonKE full memory", "G-RAG":"Qwen retrieved", "G-ALL":"Qwen full memory"}
    rows = []
    for key in ("R0", "R-RAG", "R-ALL", "G-RAG", "G-ALL"):
        m = original["conditions"][key]["overall"]; c = m["counts"]
        rows.append([labels[key], f"{c['correct']}/{m['assigned']}", f"{100*m['accuracy']:.2f}%", c["wrong"], c["abstain"]])
    report.table("All original main outcomes", ["Configuration", "Correct", "Accuracy", "Wrong", "Abstain"], rows, [2.42, 1.1, 1.1, .9, 1.01])
    report.prose("""
    All five conditions completed 256 assigned questions. No execution error or unfinished request
    remained. Both separate oracle conditions scored 16/16. The original final total is five times
    256 plus twice sixteen, or 1312 generations. The table retains the stale reference and all
    unsuccessful updated outcomes rather than presenting only the strongest method.
    """)
    report.figure(OPS / "derived/outcome-composition.png", "Original study outcome composition with 256 assigned questions per main condition. R0 is a stale diagnostic. All updated conditions receive the approved correction.")
    report.prose("""
    ReasonKE's selected updated context produced 256/256 correct answers, compared with 231/256
    under full memory. The same model therefore did not maintain a perfect result when given more
    competing facts. Qwen retrieved and full memory each scored 241/256, but with different outcome
    types. This shows why the combination of model and supplied context matters more than a single
    headline accuracy. It does not identify a causal internal attention mechanism.
    """)
    report.h2("Locality and preservation")
    report.table("Original locality correctness and preservation", ["Configuration", "All locality correct", "Baseline correct retained"], [
        ["ReasonKE initial facts", "124/128", "Reference set of 124"],
        ["ReasonKE retrieved updated facts", "128/128", "124/124"],
        ["ReasonKE full updated memory", "119/128", "117/124"]
    ], [2.7, 1.83, 2.0])
    report.prose("""
    Retaining an originally wrong answer is not factual success. The retrieved condition preserved all
    124 baseline correct locality answers and answered all 128 locality questions correctly. The full
    memory condition retained 117/124 and answered 119/128 correctly. Endorsement of an inapplicable
    edited value occurred in zero of 128 retrieved locality opportunities, three of 128 ReasonKE full
    memory opportunities and one of 128 Qwen full memory opportunities.
    """)
    report.h2("Paired original differences")
    report.figure(OPS / "derived/paired-differences.png", "Paired original accuracy differences with 95 percent intervals from resampling 32 whole lineages. Negative values mean the first listed condition scored lower.")
    report.prose("""
    ReasonKE retrieved minus Qwen retrieved was +5.859375 percentage points, with a paired interval
    from +3.90625 to +8.203125. ReasonKE full memory minus Qwen full memory was -3.90625 points,
    with an interval from -7.8125 to -0.390625. ReasonKE therefore did not beat the comparison model
    in every configuration. These are matched checkpoint comparisons, not a verified causal ablation
    of the upstream training procedure.
    """)
    report.h2("Separate depth challenge")
    report.table("Depth main results by group", ["Question group", "ReasonKE correct", "Qwen correct", "Qwen abstain"], [
        ["Locality controls", "64/64", "64/64", "0"], ["One hop", "64/64", "64/64", "0"],
        ["Two hops", "64/64", "57/64", "7"], ["Three hops", "64/64", "63/64", "1"],
        ["All main questions", "256/256", "248/256", "8"]
    ], [2.22, 1.45, 1.45, 1.41])
    report.figure(DEPTH / "derived/depth-accuracy.png", "Separate depth accuracy by relation count. Every group contains 64 assigned questions per model. Zero denotes locality controls requiring an independent fact.")
    report.prose("""
    Each checkpoint completed 272 requests: 256 main plus sixteen preselected three hop oracle
    requests. Main and oracle questions must not be pooled into a 272 question headline. ReasonKE
    main accuracy was 100 percent and Qwen main accuracy 96.875 percent. Each oracle diagnostic
    scored 16/16. Both final ledgers completed without execution errors or unfinished assignments.

    The main ReasonKE minus Qwen difference was +3.125 percentage points, with a 95 percent
    lineage bootstrap interval of +1.5625 to +5.078125. For three hops the difference was +1.5625
    points, with an interval of 0 to +4.6875. The three hop interval includes zero, so this small subset
    does not establish a stable nonzero advantage. The two studies have different question populations
    and selectors; a lower or higher cross study percentage cannot isolate the effect of hop count.
    """)
    report.h2("Support selection and answer production")
    report.figure(DEPTH / "derived/retrieval-and-answer.png", "Required support versus final answer success in the depth study. The same selector supplies complete paths for both models; Qwen still abstains on eight main questions.")
    report.prose("""
    The depth selector supplied all 448 required fact occurrences across 256 main paths for each model,
    including 192 required occurrences across all 64 three hop paths. These repeated counts describe
    one shared retrieval policy, not two independent retrieval experiments. The original selector
    supplied all 320 required occurrences and all 64 two hop paths in its separate population.

    Complete support cannot be counted as a correct neural answer. All eight depth abstentions had
    their required facts supplied. The observed failures occur after successful support selection under
    the fixed generation configuration. None of those eight question IDs belongs to the preselected
    oracle subset; the 16/16 oracle result cannot be described as repairing those failures.
    """)
    report.h2("Latency and resource cost")
    report.table("Warm selected context latency", ["Study and model", "Generated n", "p50 seconds", "p95 seconds"], [
        ["Original ReasonKE", "256", "15.06", "16.71"], ["Original Qwen", "256", "0.92", "1.20"],
        ["Depth ReasonKE", "256", "15.94", "19.28"], ["Depth Qwen", "256", "1.12", "1.38"]
    ], [2.41, 1.15, 1.48, 1.49])
    report.figure(DEPTH / "derived/latency.png", "Depth pipeline latency distributions. The main conditions each have 256 generated responses; oracle conditions each have sixteen. Abstentions remain in the measured latency population.")
    report.prose("""
    Warm request time includes context selection, local process communication and generation. Loading,
    downloads, browser network and operator actions are excluded. Qwen's times include its abstained
    outputs. ReasonKE produced more correct main answers in these samples while taking materially
    longer. A deployment decision must account for this delay; accuracy alone is an incomplete
    service metric. The approximately 6.59 GiB loading allocation also exceeded the laptop's roughly
    6 GiB GPU capacity for this fixed loader, so a T4 was used. This does not prove that every possible
    quantized 7B implementation is impossible on that laptop.
    """)

    report.chapter("9 Failure analysis and the perfect score")
    report.h2("Retained failures across configurations")
    report.prose("""
    The original study provides important counterevidence to a universal reliability claim. ReasonKE
    full memory produced 25 strict wrong answers although required support was present. Inspection
    found that each endorsed an answer belonging to another supplied fact. Qwen retrieved produced
    fifteen literal UNKNOWN abstentions, all among its 32 two hop paraphrases. These remain zero
    credit outcomes even though Qwen's answered only precision was perfect.

    Qwen full memory retained thirteen strict wrong answers and two abstentions. Not every strict
    mismatch is a false factual endorsement: two outputs include a compatible owner and contact chain
    but violate the person only answer contract, and another combines a Pro owner with the requested
    contact. We retain the fixed scores while explaining that distinction. Changing aliases after test
    inspection would invalidate the frozen evaluation.
    """)
    report.table("All eight depth abstentions", ["Lineage", "Category", "Expected", "Actual"], [
        ["depth-001", "Two hop", "Diya Desai", "UNKNOWN"],
        ["depth-004", "Two hop", "Arjun Desai", "UNKNOWN"],
        ["depth-016", "Two hop", "Arjun Nair", "UNKNOWN"],
        ["depth-024", "Three hop paraphrase", "Indore", "UNKNOWN"],
        ["depth-025", "Two hop", "Diya Kapoor", "UNKNOWN"],
        ["depth-028", "Two hop", "Arjun Kapoor", "UNKNOWN"],
        ["depth-031", "Two hop", "Ishita Kapoor", "UNKNOWN"],
        ["depth-040", "Two hop", "Arjun Menon", "UNKNOWN"]
    ], [1.15, 2.22, 1.85, 1.31])
    report.prose("""
    The seven direct two hop abstentions concentrate in one final wording template: identify the
    service's release 2 owning team and give only its on call contact's name. Qwen abstained on seven
    of nine instantiations, whereas the other direct two hop templates scored 23/23. All 32 two hop
    paraphrases were correct. This descriptive concentration does not prove an internal cause.
    It also explains why the observed three hop group scored higher than the two hop group; a
    monotonic decline with depth is not supported by this sample.
    """)
    report.h2("What the 100 percent result means")
    report.prose("""
    The correct claim is observed controlled benchmark accuracy. ReasonKE returned an accepted
    final short answer for all 256 assigned main questions in each of two separately frozen studies.
    This is a genuine recorded result, not a target filled from a database or an upstream paper score.
    It should neither be hidden nor described as a promise about arbitrary questions.

    Several properties make ceiling performance plausible. Each world contains just nine or twelve
    explicit facts. The tested paths are short and the relation families repeat. The selector supplies
    every required fact in these final populations. Answers are short canonical names or places, and
    wording templates repeat within each split. Even in the depth study, new entity assignments and
    separate wording banks remain close to the same constructed task family. The test has limited
    power to distinguish two strong models on its easiest groups.

    There are 32 held out lineages per study, each producing eight correlated questions. The 256
    variants are not 256 independent arbitrary conversations. Automatically generated or authored
    labels were mechanically audited, not independently reviewed by people. Hash checks demonstrate
    evidence integrity and identity; they do not prove semantic validity, eliminate every form of
    leakage or make the catalogue representative of a real organisation.

    The perfect score therefore establishes that the chosen system completed this specified contract
    under its fixed inputs. It does not establish reliable language ambiguity detection, open domain
    retrieval, many simultaneous unrelated edits, multilingual use, adversarial sources or correctness
    of an unrestricted conversation. A bootstrap can have zero width when every sampled lineage
    succeeds; that arithmetic does not imply certainty on future traffic. We do not add invented
    failures or alter the test to make a percentage appear more realistic.
    """)
    report.h2("Answers and explanations are different claims")
    report.prose("""
    A separate original explanation audit selected 32 traces before inspection, sixteen from ReasonKE
    retrieved and sixteen from full memory. For the retrieved traces all 32 extracted indexed
    subanswers were supported. Full memory had 26 supported, three contradicted and three unsupported
    indexed claims. Each condition also contained 64 prose segments that the bounded audit did not
    assess. No complete trace was certified as fully verified. A correct final name can coexist with
    unsupported explanation text, so the interface displays supplied provenance without certifying
    every generated sentence.
    """)
    report.h2("Threats to validity and the next useful study")
    report.prose("""
    Internal validity is supported by fixed questions, matched settings, all outcome denominators and
    replayable raw outputs, but strict formatting rules can disagree with a human assessment of useful
    information. Construct validity is limited by template based questions and narrow graph relations.
    External validity is limited by small fictional memories and absence of real user traffic.
    Checkpoint comparisons do not isolate a causal training effect because exact parent ancestry is
    not verified. The original and depth populations cannot be pooled into a single improved score.

    A future study should first freeze independently human reviewed questions from realistic operational
    records, with broader wording, genuinely ambiguous requests, missing evidence and larger distractor
    collections. It should measure safe refusal and error severity as well as exact answers and latency.
    Such a study would test the limits suggested by these results. It is future validation, not a
    retroactive modification of the completed benchmark or a condition for reporting its observations.
    """)

    report.chapter("10 Application workflow and verification")
    report.h2("Using the website")
    report.prose("""
    Chat accepts ordinary text conversations and factual questions. Knowledge contains the correction
    library, explicit approval editor, canonical IDs, source metadata, version history and rollback.
    Research results displays the original text experiment and separate depth challenge, with raw
    evidence exports and scientific graphs. The current ReasonKE worker is text only; the image
    capability of an archived backend is not part of this final model evaluation.

    In a fresh workspace, approve the nine fictional India facts. Ask who owns Kaveri Payment Gateway
    in release 2 and ask for the contact's office city. Draft the provided ownership update and explicitly
    approve the replacement. Ask the direct and three hop questions again; ask the hosting and release 1
    questions to inspect locality. Use Compare original and corrected with approved memory revision 9
    to generate historical and current answers. Roll back the replacement and ask the same questions
    anew. A previously stored chat answer stays unchanged; the new response uses restored memory.

    Revision 9 identifies the approved starting handbook, revision 10 the ownership replacement and
    revision 11 its atomic rollback. These numbers refer to this finite demonstration, not to every
    future user workspace. Normal conversation remains available, but this factual benchmark does not
    score the quality or truth of every possible conversational reply.
    """)
    report.h2("Measured website workload")
    report.table("India application acceptance populations", ["Stage", "Assigned neural answers"], [
        ["Initial direct two hop and three hop", "3"],
        ["Edited paraphrase paths and locality", "6"],
        ["Historical and current office comparison", "2"],
        ["Rollback direct two hop and three hop", "3"],
        ["Total actual model generations", "14"]
    ], [4.6, 1.93])
    report.prose("""
    All fourteen generated answers matched their fixed short answer contracts. All twelve recorded
    provenance and transaction consistency checks passed, including rejected competing approval,
    unchanged memory after rejection, historical comparison and restoration of the initial nine active
    facts. Thirteen HTTP generation requests produced fourteen answers because one comparison invoked
    the model twice. Request p50 and p95 were 17.11 and 33.39 seconds, with 268.45 seconds total
    workload time. The mixed request quantiles are not per generation latency.

    The laptop replay verified raw answer extraction, exact model revision, iterative context policy,
    expected memory revisions, all scores and the saved rollback contents without new model calls.
    This is one fictional display lineage, not fourteen independent held out facts. The earlier website
    supplement also completed twelve workflow and four history answers, with eleven recorded
    consistency checks. Neither operational workload changes a frozen benchmark denominator.
    """)
    report.h2("Deterministic and neural evidence")
    report.prose("""
    Transaction and deployment tests can establish that a conflict is rejected or a memory snapshot
    restores correctly. They cannot establish that the model generated the correct contact. The
    engineering deployment receipt records 24 software checks with zero model generations and null
    model accuracy. Conversely, a model's correct answer does not prove atomic database behaviour.
    The project retains both types of evidence with separate labels.
    """)

    report.chapter("11 Engineering and reproducibility")
    report.h2("Minimal deployment stack")
    report.prose("""
    The browser uses HTML, CSS and JavaScript. The website uses Python's standard library HTTP server
    and SQLite, with a small pinned image dependency for the preserved application surface. The private
    neural worker adds PyTorch, Transformers, Accelerate, bitsandbytes and safetensors in an isolated
    environment. The lightweight application container excludes weights, private chat databases and
    session secrets. It can serve saved evidence without a GPU, but requires a ready model worker for
    live neural answers. It is not a self contained GPU image.

    The local server binds loopback, validates Host and exact Origin and requires a session token for
    changes. The worker uses a separate bearer credential and rejects direct browser origin requests.
    Colab permits the exact configured HTTPS proxy origin. This is a single operator research
    deployment. It does not provide enterprise role based access control, tenant isolation, distributed
    consistency or a security certification.
    """)
    report.h2("Recovery and evidence integrity")
    report.prose("""
    Final protocols bind source files, data, settings, scoring and completed development inputs by
    SHA256. Raw ledgers retain task identity, selected facts, raw model text, extracted answer, outcome,
    tokens, latency and resources. CPU replay reconstructs the request and scores the saved text.
    Figure manifests bind scientific exports to those results. A checksum detects mismatched bytes
    relative to the saved manifest; it is not proof of factual truth or independent authorship.

    An earlier temporary development runtime expired before outputs were downloaded. Those unknown
    outcomes are disclosed and excluded from measured evidence. Saved profiles were restored byte
    for byte for a replacement development run before the depth protocol was frozen. A later browser
    disconnection left the final evaluator alive. Preservation checks confirm that all previously saved
    174, 197, 235 and 247 ReasonKE records stayed unchanged in the completed 272 record ledger.
    No completed failure was rerun to improve a score.

    Portable text workspace exports retain conversations, facts, versions and events. Restore validates
    schema, bounded size, references, lineage consistency and agreement of event replay with state,
    then imports transactionally into an empty workspace. An export is private user data; it is separate
    from the public fictional benchmark receipts. A temporary Colab session can end without losing
    published results, but live chat needs the pinned worker loaded again.
    """)
    report.h2("Regression and continuous integration")
    report.prose("""
    Automated checks cover fact identity and lifecycle transitions, atomic rollback, historical reads,
    conversation history, request validation, worker contracts, dataset splits, retrieval, scoring,
    evidence tampering, recovery, figure export and the finite India workflow audit. JavaScript syntax,
    container construction and HTTP deployment checks run in GitHub Actions. The completed release
    ae7bb2010a82fd1d744d008eb7c6f870198aece5 passed all six workflow jobs [10]. These software checks
    provide regression evidence; their count does not increase model accuracy.
    """)

    report.chapter("12 Conclusion and future work")
    report.prose("""
    EditGuard implements a usable maintenance workflow around an attributed learned generator. An
    operator can approve a scoped factual replacement, inspect its provenance, compare historical and
    current answers, reject an unresolved competing approval and restore the previous state atomically.
    The independent datasets, fair updated contexts, finite runners and raw evidence audits make the
    resulting behaviour measurable and reviewable.

    The completed original and depth studies support a bounded finding: selected approved facts enabled
    ReasonKE to produce every requested short main answer correctly in these controlled packets, while
    the matched Qwen conditions retained abstentions. The same ReasonKE checkpoint failed under full
    memory and had substantially higher latency. The three hop subset difference interval includes
    zero, and explanation text is not fully verified. These findings support the integrated design without
    claiming that one model or editing method is universally superior.

    The project delivers an engineered and evaluated platform around an upstream trained checkpoint.
    It does not claim new foundation model training, faithful SERAC replication or guaranteed dialogue
    accuracy. The next validation should use independently reviewed operational questions, larger
    memories, ambiguous scopes and missing facts.
    """)

    report.chapter("References")
    refs = [
        "[1] P Lewis and colleagues. Retrieval Augmented Generation for Knowledge Intensive NLP Tasks. NeurIPS 2020. https://arxiv.org/abs/2005.11401",
        "[2] E Mitchell C Lin A Bosselut C D Manning and C Finn. Memory Based Model Editing at Scale. ICML 2022. https://arxiv.org/abs/2206.06520",
        "[3] K Meng D Bau A Andonian and Y Belinkov. Locating and Editing Factual Associations in GPT. NeurIPS 2022. https://arxiv.org/abs/2202.05262",
        "[4] Y Wu L Ding L Shen and D Tao. Robust Knowledge Editing via Explicit Reasoning Chains for Distractor Resilient Multi Hop QA. Official Reason KE repository and citation. https://github.com/YukinoshitaKaren/Reason-KE",
        "[5] Y Wu L Ding L Shen and D Tao. Reason KE++ Aligning the Process Not Just the Outcome for Faithful LLM Knowledge Editing. 2025. https://arxiv.org/abs/2511.12661",
        "[6] YukinoKaren. ReasonKE_plus author published checkpoint. Revision ceaedc2ae7c8fed6257953587d991ce550db6106. https://huggingface.co/YukinoKaren/ReasonKE_plus . Exact checkpoint to paper score association and weight license metadata remain undocumented in the inspected evidence.",
        "[7] Qwen team. Qwen2.5 7B Instruct official model card. Revision a09a35458c702b33eeacc393d103063234e8bc28. https://huggingface.co/Qwen/Qwen2.5-7B-Instruct/blob/a09a35458c702b33eeacc393d103063234e8bc28/README.md",
        "[8] EditGuard project. Completed original operations evidence and frozen protocol. reports/reasonke/operations-final . https://github.com/ShivamRaghav7/verity/tree/ae7bb2010a82fd1d744d008eb7c6f870198aece5/reports/reasonke/operations-final",
        "[9] EditGuard project. Completed separate depth evidence and India application audit. reports/reasonke/depth-final . https://github.com/ShivamRaghav7/verity/tree/ae7bb2010a82fd1d744d008eb7c6f870198aece5/reports/reasonke/depth-final",
        "[10] EditGuard project. Passing release regression workflow. 28 October 2026. https://github.com/ShivamRaghav7/verity/actions/runs/40723337854"
    ]
    for item in refs: report.p(item)
    report.prose("""
    Primary research describes the method families; local raw receipts establish the numerical findings
    reported here. Published upstream scores are not our results. The final deployment differs from the
    earlier archived auxiliary model and visual studies. Their histories are retained in the repository
    but excluded from the two current scorecards. Web references were checked on 28 October 2026.
    """)

    report.chapter("Appendix A Reproduction and evidence map")
    report.h2("Released evidence")
    report.table("Evidence artifacts and their purpose", ["Repository location", "Purpose"], [
        ["reports/reasonke/operations-final", "Original raw outputs metrics protocol explanation audit and failures"],
        ["reports/reasonke/depth-final", "Depth raw outputs metrics protocol failures and preservation audit"],
        ["depth-final/india-acceptance.json", "All fourteen neural application responses and transaction snapshots"],
        ["depth-final/india-host-audit.json", "CPU replay of the application receipt with zero model calls"],
        ["Each study derived directory", "CSV tables PNG and SVG figures and source hash manifests"],
        ["docs/INDIA_DEMO.md", "Operator walkthrough including comparison rollback and private export"],
        ["docs/CHAT_API.md", "Application API reference"],
        ["notebooks/07_india_chat.ipynb", "Pinned India website launcher without benchmark execution"],
        ["notebooks/06_depth_challenge.ipynb", "Staged fixed profile development freeze and final experiment"]
    ], [3.18, 3.35])
    report.h2("Recompute scores without another model call")
    report.p("Run these from the repository using an isolated Python environment. Use fresh output paths; preserve the released evidence. The commands below audit saved responses and do not generate new answers.")
    report.code("""python -m editguard_depth.metrics \\
  reports/reasonke/depth-final/final-reasonke.json \\
  reports/reasonke/depth-final/final-instruction.json \\
  --protocol reports/reasonke/depth-final/protocol.json \\
  --output artifacts/depth-replay-results.json

python -m scripts.audit_india_http \\
  reports/reasonke/depth-final/india-acceptance.json \\
  --output artifacts/india-replay-audit.json""")
    report.p("These are shell continuation examples. In PowerShell, put the command on one line or use its backtick continuation syntax. Do not copy the backslash as a PowerShell continuation.")
    report.h2("Pinned identities")
    for label, value in [
        ("ReasonKE model revision", "ceaedc2ae7c8fed6257953587d991ce550db6106"),
        ("Qwen model revision", "a09a35458c702b33eeacc393d103063234e8bc28"),
        ("Depth inference source", "7da6d1df5a790c51c7cc95d4378117f5eae052c3"),
        ("Depth protocol SHA256", sha(DEPTH / "protocol.json")),
        ("Original protocol SHA256", sha(OPS / "protocol.json"))]:
        label_p = report.p(label, bold=True)
        label_p.paragraph_format.keep_with_next = True
        report.code(value)
    report.p("The depth protocol was published in commit 08d274e204e0dc8c7fcff94cc3aabcefda0debef before final inference. Exact source and data identities are in each protocol. Complete source restoration, rather than only copying a result JSON, is required to rerun an experiment.")

    report.chapter("Appendix B Metric glossary and worked example")
    report.table("How to read the report numbers", ["Term", "Interpretation"], [
        ["256 main questions", "Eight variants per 32 held out graph lineages in one study and configuration"],
        ["272 requests per depth model", "256 main plus sixteen separate oracle diagnostics"],
        ["Primary accuracy", "Correct divided by all assigned questions including unsuccessful outcomes"],
        ["Coverage", "Fraction producing an endorsed answer rather than abstaining or failing"],
        ["Retention", "Previously correct locality answers that remain correct"],
        ["Required fact occurrence", "A necessary supporting fact counted once for each question needing it"],
        ["Complete path", "Every required supporting fact was supplied for the question"],
        ["p50 and p95", "Median and 95th percentile of the stated measured latency population"],
        ["Percentage point difference", "Direct subtraction of two percentages not relative percent improvement"],
        ["Paired lineage interval", "Bootstrap uncertainty with complete graph lineages resampled together"],
        ["Oracle support", "Gold supporting facts for a preselected diagnostic not normal retrieval"],
        ["Revision", "Approved memory state distinct from a model weight checkpoint"]
    ], [1.95, 4.58])
    report.prose("""
    For the depth main comparison, Qwen produced 248 accepted answers and eight UNKNOWN responses.
    Primary accuracy is 248 divided by 256, or 96.875 percent. Answered only precision is 248 divided
    by 248, or 100 percent. Reporting only the latter would hide the abstentions. ReasonKE's 256
    divided by 256 gives an observed 100 percent, so the direct difference is 3.125 percentage points.

    For original locality, the initial model answered 124 of 128 controls correctly. Preserving those
    124 answers gives 124/124 retention, while answering all controls correctly gives 128/128 factual
    locality accuracy. They answer different questions and require different denominators. Neither
    number proves that all possible unrelated questions are preserved.
    """)
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    d.save(output)
    manifest = {"schema": "editguard-complete-docx-report-v1", "created_utc": datetime.now(timezone.utc).isoformat(),
        "output": str(output.relative_to(ROOT)), "sha256_before_layout_refresh": sha(output), "model_calls": 0,
        "builder_sha256": sha(__file__), "inputs": {str(p.relative_to(ROOT)):sha(p) for p in [
            OPS/"results.json", OPS/"protocol.json", DEPTH/"results.json", DEPTH/"protocol.json",
            DEPTH/"india-acceptance.json", DEPTH/"india-host-audit.json"]},
        "figures": report.figures, "tables": report.tables,
        "scope": "Final ReasonKE platform with original and separate depth studies; no invented project training losses"}
    (output.parent / "report-build-manifest.json").write_text(json.dumps(manifest, indent=2)+"\n", encoding="utf-8")
    print(json.dumps({"output":str(output), "figures":len(report.figures), "tables":report.tables, "model_calls":0}))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=ROOT/"output/docx/EditGuard_Project_Report.docx")
    build(parser.parse_args().output)
